"""Banco maestro — motor real, no el Excel legacy (autorización del Founder,
19-sep-2026, ver docstring de `banco_maestro.py`). Casos normal/límite/
adversarial sobre generación, anti-repetición, ENTREGADO y exportación."""

import csv
import os
import tempfile
import unittest

import banco_maestro as bm
import editorial
import universe
from semantic_fingerprint import SemanticFingerprint
from semantic_memory import ENTREGADO, SemanticMemory


class TestEstadoEntregadoEnSemanticMemory(unittest.TestCase):
    """`ENTREGADO` es el estado que pide el requisito 4 del mandato: "una vez
    utilizado, regístralo como ENTREGADO y bloquéalo para futuras
    generaciones". Se prueba directamente sobre `semantic_memory.py`, no
    sólo indirectamente a través de `banco_maestro.py`."""

    def test_entregado_esta_en_los_estados_conocidos(self):
        import semantic_memory as sm
        self.assertIn(ENTREGADO, sm.ESTADOS)

    def test_entregado_no_es_memoria_fuerte(self):
        """Entregarse en un banco no es que el Founder lo haya aprobado —
        mezclarlo con MEMORIA_FUERTE sesgaría `preferencias()` con volumen
        de producción, no con gusto real."""
        import semantic_memory as sm
        self.assertNotIn(ENTREGADO, sm.MEMORIA_FUERTE)

    def test_entregado_bloquea_un_equivalente_semantico(self):
        memoria = SemanticMemory()
        fp = SemanticFingerprint(content_id="A", materia="civil", concepto_nucleo="x",
                                 familia_editorial="mito", necesidad="entender")
        memoria.record(fp, ENTREGADO, "lote-1")
        parafraseado = SemanticFingerprint(content_id="B", materia="civil",
                                           concepto_nucleo="x", familia_editorial="mito",
                                           necesidad="entender", hook="otro hook",
                                           formato="otro formato")
        veredicto = memoria.evaluar(parafraseado)
        self.assertTrue(veredicto.bloquea)
        self.assertIn("ENTREGADO", veredicto.motivo)

    def test_entregado_tiene_ventana_permanente_como_historica(self):
        import semantic_memory as sm
        self.assertEqual(sm.COOLDOWN[ENTREGADO], sm.COOLDOWN[sm.HISTORICA])


class TestMapaDeCategoriasFounder(unittest.TestCase):
    """Requisito 2: mitos jurídicos, diferencias, conceptos, pasos
    prácticos, errores frecuentes, derechos, obligaciones, casos, y otros
    formatos ya definidos. Todas las familias que se citan deben existir de
    verdad en el registro (CLAUDE.md §2: no asumir capacidades)."""

    def setUp(self):
        self.universo = editorial.EditorialUniverse.load()

    def test_las_8_categorias_estan_presentes(self):
        esperadas = {"mitos_juridicos", "diferencias", "conceptos", "pasos_practicos",
                    "errores_frecuentes", "derechos", "obligaciones", "casos"}
        self.assertEqual(esperadas, set(bm.CATEGORIAS_FOUNDER))

    def test_todas_las_familias_citadas_existen_en_el_registro_real(self):
        nombres_reales = set(self.universo.names())
        for categoria, familias in bm.CATEGORIAS_FOUNDER.items():
            for f in familias:
                self.assertIn(f, nombres_reales,
                              f"{categoria} cita {f!r}, que no existe en el registro real.")

    def test_familia_a_categoria_reconoce_las_8_nombradas(self):
        self.assertEqual(bm.familia_a_categoria("mito"), "mitos_juridicos")
        self.assertEqual(bm.familia_a_categoria("diferencia"), "diferencias")
        self.assertEqual(bm.familia_a_categoria("derecho"), "derechos")
        self.assertEqual(bm.familia_a_categoria("obligacion"), "obligaciones")

    def test_familia_no_nombrada_cae_en_otros_formatos(self):
        self.assertEqual(bm.familia_a_categoria("quiz_juridico"), "otros_formatos_ya_definidos")

    def test_grupos_de_generacion_cubren_las_65_familias_sin_perder_ninguna(self):
        grupos = bm._grupos_de_generacion(self.universo)
        cubiertas = set()
        for _, familias in grupos:
            cubiertas |= set(familias)
        self.assertEqual(cubiertas, set(self.universo.names()))


class TestGeneracionCategorizada(unittest.TestCase):
    def setUp(self):
        self.universo = editorial.EditorialUniverse.load()
        _, self.materias = universe.cargar_materias()

    def test_reserva_categorizada_cubre_las_9_puertas_editoriales(self):
        import random
        rng = random.Random(42)
        candidatos = bm.generar_reserva_categorizada(rng, self.universo, self.materias,
                                                      objetivo=90, factor_reserva=1)
        categorias_vistas = {bm.familia_a_categoria(c.familia_editorial) for c in candidatos}
        esperadas = set(bm.CATEGORIAS_FOUNDER) | {"otros_formatos_ya_definidos"}
        self.assertEqual(categorias_vistas, esperadas)

    def test_reserva_categorizada_reparte_proporcionalmente(self):
        """Con un total grande y 9 grupos en round-robin, ningún grupo
        debería quedar con menos de la mitad del reparto uniforme esperado
        -- el hallazgo real que este módulo corrige es que una familia
        nombrada por el Founder quedara casi sin representar."""
        import random
        rng = random.Random(7)
        objetivo = 180
        candidatos = bm.generar_reserva_categorizada(rng, self.universo, self.materias,
                                                      objetivo=objetivo, factor_reserva=1)
        conteo = {}
        for c in candidatos:
            cat = bm.familia_a_categoria(c.familia_editorial)
            conteo[cat] = conteo.get(cat, 0) + 1
        esperado_uniforme = objetivo / 9
        for cat, n in conteo.items():
            self.assertGreaterEqual(n, esperado_uniforme * 0.5,
                                    f"{cat} quedó infrarrepresentada: {n}")

    def test_candidatos_siguen_no_verificados(self):
        """CLAUDE.md §4: ninguna IA es fuente jurídica. Todo candidato de
        este banco tiene que seguir apuntando a verificación humana."""
        import random
        rng = random.Random(1)
        candidatos = bm.generar_reserva_categorizada(rng, self.universo, self.materias,
                                                      objetivo=10, factor_reserva=1)
        for c in candidatos:
            self.assertEqual(c.estado_verificacion, "NO_VERIFICADO")
            self.assertEqual(c.proxima_accion, "legalmente-legal-verification")


class TestConstruirBancoCasoNormal(unittest.TestCase):
    def test_produce_el_objetivo_pedido_sin_duplicados_internos(self):
        reporte = bm.construir_banco(objetivo=25, seed=555, factor_reserva=3, persistir=False)
        self.assertEqual(len(reporte.entregados), 25)
        fps = [c.fingerprint() for c in reporte.entregados]
        for i, a in enumerate(fps):
            for b in fps[i + 1:]:
                self.assertFalse(a.equivalente_a(b),
                                 f"{a.content_id} y {b.content_id} son equivalentes: "
                                 "el mismo tema con otra ropa no debería pasar (requisito 5).")

    def test_ningun_entregado_repite_el_historial_real(self):
        """Requisito 3: comparar contra el historial completo ANTES de
        incorporar. Se re-verifica con una memoria fresca, independiente de
        la que usó `construir_banco` internamente."""
        reporte = bm.construir_banco(objetivo=20, seed=777, factor_reserva=3, persistir=False)
        historial = bm.cargar_memoria_historial_completo()
        for c in reporte.entregados:
            veredicto = historial.evaluar(c.fingerprint())
            self.assertFalse(veredicto.bloquea,
                             f"{c.candidate_id} repite historial real: {veredicto.motivo}")

    def test_clasificacion_disponible_utilizado_bloqueado_es_disjunta(self):
        reporte = bm.construir_banco(objetivo=15, seed=321, factor_reserva=4, persistir=False)
        ids_entregados = {c.candidate_id for c in reporte.entregados}
        ids_disponibles = {c.candidate_id for c in reporte.disponibles}
        ids_bloqueados = {cid for cid, _ in reporte.bloqueados}
        self.assertEqual(set(), ids_entregados & ids_disponibles)
        self.assertEqual(set(), ids_entregados & ids_bloqueados)
        self.assertEqual(set(), ids_disponibles & ids_bloqueados)

    def test_variedad_real_de_materias_y_familias_en_el_lote_entregado(self):
        """Requisito 6: variedad real de materias jurídicas y ángulos."""
        reporte = bm.construir_banco(objetivo=40, seed=909, factor_reserva=3, persistir=False)
        materias = {c.materia for c in reporte.entregados}
        familias = {c.familia_editorial for c in reporte.entregados}
        self.assertGreaterEqual(len(materias), 8)
        self.assertGreaterEqual(len(familias), 8)


class TestPersistenciaYBloqueoFuturo(unittest.TestCase):
    """Requisito 4: "una vez utilizado, regístralo como ENTREGADO y
    bloquéalo para futuras generaciones" — la prueba central de este
    módulo."""

    def test_una_segunda_corrida_con_la_misma_semilla_no_repite_lo_ya_entregado(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "banco.json")
            r1 = bm.construir_banco(objetivo=15, seed=2024, factor_reserva=3,
                                    banco_previo_path=path, persistir=True)
            r2 = bm.construir_banco(objetivo=15, seed=2024, factor_reserva=3,
                                    banco_previo_path=path, persistir=True)
            ids1 = {c.candidate_id for c in r1.entregados}
            ids2 = {c.candidate_id for c in r2.entregados}
            self.assertEqual(set(), ids1 & ids2,
                             "la segunda corrida reentregó exactamente lo mismo que la primera.")
            fp1 = [c.fingerprint() for c in r1.entregados]
            fp2 = [c.fingerprint() for c in r2.entregados]
            equivalentes = [(a.content_id, b.content_id) for a in fp2 for b in fp1
                            if a.equivalente_a(b)]
            self.assertEqual([], equivalentes,
                             "la segunda corrida entregó algo semánticamente equivalente a la "
                             "primera aunque los IDs fueran distintos.")

    def test_los_bloqueados_de_la_segunda_corrida_citan_entregado(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "banco.json")
            bm.construir_banco(objetivo=10, seed=333, factor_reserva=3,
                               banco_previo_path=path, persistir=True)
            r2 = bm.construir_banco(objetivo=10, seed=333, factor_reserva=3,
                                    banco_previo_path=path, persistir=True)
            motivos_entregado = [m for _, m in r2.bloqueados if "ya ENTREGADO" in m]
            self.assertGreater(len(motivos_entregado), 0,
                               "la segunda corrida con la misma semilla debería reencontrar y "
                               "bloquear candidatos ya ENTREGADO en la primera.")

    def test_sin_persistir_no_deja_rastro_en_disco(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "banco.json")
            bm.construir_banco(objetivo=5, seed=1, factor_reserva=3,
                               banco_previo_path=path, persistir=False)
            self.assertFalse(os.path.exists(path))

    def test_banco_previo_vacio_si_no_hay_archivo(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "no-existe.json")
            previo = bm.cargar_banco_previo(path)
            self.assertEqual(0, len(previo))


class TestExportacion(unittest.TestCase):
    def test_exportar_csv_tiene_las_columnas_y_filas_esperadas(self):
        reporte = bm.construir_banco(objetivo=8, seed=44, factor_reserva=3, persistir=False)
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "banco.csv")
            bm.exportar_csv(reporte.entregados, path)
            with open(path, newline="", encoding="utf-8") as f:
                filas = list(csv.reader(f))
            self.assertEqual(filas[0], bm.COLUMNAS)
            self.assertEqual(len(filas) - 1, len(reporte.entregados))
            self.assertIn("ENTREGADO", filas[1])

    def test_exportar_excel_produce_tres_hojas_clasificadas(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no disponible en este entorno")
        reporte = bm.construir_banco(objetivo=6, seed=66, factor_reserva=4, persistir=False)
        with tempfile.TemporaryDirectory() as d:
            excel_path = os.path.join(d, "banco.xlsx")
            csv_path = os.path.join(d, "banco.csv")
            rutas = bm.exportar_banco(reporte, excel_path=excel_path, csv_path=csv_path)
            wb = openpyxl.load_workbook(rutas["excel"])
            self.assertEqual(set(wb.sheetnames), {"ENTREGADO", "DISPONIBLE", "BLOQUEADO"})
            ws = wb["ENTREGADO"]
            self.assertEqual(ws.max_row - 1, len(reporte.entregados))
            self.assertEqual([c.value for c in ws[1]], bm.COLUMNAS)

    def test_exportar_excel_sin_openpyxl_lanza_error_claro(self):
        original = bm.openpyxl
        try:
            bm.openpyxl = None
            with self.assertRaises(bm.BancoMaestroError):
                bm.exportar_excel([], path="/tmp/no-deberia-crearse.xlsx")
        finally:
            bm.openpyxl = original


class TestAdversarial(unittest.TestCase):
    def test_categorias_founder_con_familia_inexistente_falla_alto_y_claro(self):
        original = dict(bm.CATEGORIAS_FOUNDER)
        try:
            bm.CATEGORIAS_FOUNDER["inventada"] = ("familia-que-no-existe",)
            universo = editorial.EditorialUniverse.load()
            with self.assertRaises(bm.BancoMaestroError):
                bm._grupos_de_generacion(universo)
        finally:
            bm.CATEGORIAS_FOUNDER.clear()
            bm.CATEGORIAS_FOUNDER.update(original)

    def test_objetivo_mayor_que_la_reserva_disponible_no_rellena_con_basura(self):
        """Si la reserva no alcanza para `objetivo`, el banco debe declarar
        menos entregados en vez de inventar candidatos o aceptar
        duplicados -- misma disciplina que `universe.select_batch`."""
        reporte = bm.construir_banco(objetivo=500, seed=3, factor_reserva=1, persistir=False)
        self.assertLessEqual(len(reporte.entregados), 500)
        fps = [c.fingerprint() for c in reporte.entregados]
        for i, a in enumerate(fps):
            for b in fps[i + 1:]:
                self.assertFalse(a.equivalente_a(b))


if __name__ == "__main__":
    unittest.main()
